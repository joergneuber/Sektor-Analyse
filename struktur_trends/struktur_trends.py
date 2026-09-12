#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
NEUBER MACRO & MARKETS
Strukturelle Markttrends – eigenständiger V1-Datenbaustein

Zweck:
    Ausschließlich Rohdaten beschaffen, validieren und in EINEM Cache speichern.
    Keine makroökonomische oder marktbezogene Interpretation.

Dateien:
    struktur_trends.py
    struktur_trends_cache.json

WICHTIG:
    - Alle Quellen müssen kostenlos/public sein.
    - Keine Schätzungen durch dieses Programm.
    - Quelleneigene Schätz-/Estimate-Flags werden erhalten.
    - Bei Abruffehlern wird ein vorhandener gültiger Cache nicht zerstört.
    - Jede Datenquelle hat ihre eigene Aktualitätsprüfung.
    - Die IEA-SDMX-Mappings werden bewusst nicht erfunden. Für MESGEN/MESBAL
      werden die offiziellen Mapping-Informationen benötigt.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import re
import sys
import time
import random
import os
import tempfile
from urllib.parse import quote
import urllib.error
import urllib.request
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parent
CACHE_FILE = SCRIPT_DIR / "struktur_trends_cache.json"

CACHE_VERSION = "1.2"

# ---------------------------------------------------------------------------
# Datenquellen / feste Konfiguration
# ---------------------------------------------------------------------------

OECD_STAN_BASE = (
    "https://sdmx.oecd.org/public/rest/data/"
    "OECD.STI.PIE,DSD_STAN@DF_STAN_2025,1.0/"
)

OECD_PRODUCTIVITY_BASE = (
    "https://sdmx.oecd.org/public/rest/data/"
    "OECD.SDD.TPS,DSD_PDB@DF_PDB,2.0/"
)

OECD_PRODUCTIVITY_GROWTH_BASE = (
    "https://sdmx.oecd.org/archive/rest/data/"
    "OECD.SDD.TPS,DSD_PDB@DF_PDB_GR,1.0/"
)

OECD_STAN_REFERENCE_AREAS = [
    "USA", "DEU", "FRA", "ITA", "GBR",
    "ESP", "NLD", "SWE", "JPN", "KOR",
]

OECD_STAN_ACTIVITIES = {
    "_T": "Total economy",
    "C": "Manufacturing",
    "C20_21": "Chemicals and pharmaceuticals",
    "C26": "Computer, electronic and optical products",
    "C27": "Electrical equipment",
    "C28": "Machinery",
    "C29": "Motor vehicles",
    "J": "Information and communication",
}

OECD_STAN_MEASURES = {
    "B1G": "Value added",
    "P51G": "Gross fixed capital formation",
    # Working-hours query is handled separately because the STAN
    # dimension structure for labour input must not be guessed.
    "WORKING_HOURS": "Working hours / labour input",
}

OECD_PRODUCTIVITY_REFERENCE_AREAS = [
    "USA", "DEU", "FRA", "ITA", "GBR",
    "ESP", "NLD", "SWE", "JPN", "KOR",
]

# Official OECD productivity measure identified during source review.
OECD_PRODUCTIVITY_MEASURE = "GVAHRS"

# Stanford AI Index is an annual public-data release, not a stable daily API.
# The exact public-data download URL can change between annual editions.
STANFORD_PUBLIC_DATA_URL = "https://hai.stanford.edu/ai-index/2026-ai-index-report"
STANFORD_ECONOMY_URL = "https://hai.stanford.edu/ai-index/2026-ai-index-report/economy"
STANFORD_RD_URL = "https://hai.stanford.edu/ai-index/2026-ai-index-report/research-and-development"
STANFORD_TAKEAWAY_URL = "https://hai.stanford.edu/news/inside-the-ai-index-12-takeaways-from-the-2026-report"
STANFORD_PUBLIC_DATA_FOLDER = "https://drive.google.com/drive/folders/1zJTOg0iR0j5SijCwFutwWvDt143lW277"

IEA_API_BASE = "https://growth-sis-cc-api-wv.iea.org/rest"
IEA_API_BASE_STABLE = "https://sis-cc-api-stable.iea.org/rest"
IEA_API_BASE_NSI_STABLE = "https://sis-cc-nsi-stable.iea.org/rest"
IEA_MAPPING_URLS = {
    "MESGEN": "https://iea.blob.core.windows.net/assets/2489b143-bc40-4b36-bbb5-2e5ac60679fd/MESGENmapping.xlsx",
    "MESBAL": "https://iea.blob.core.windows.net/assets/bbc02b8a-b510-471a-8597-1899f302a57b/MESBALmapping.xlsx",
}

SIPRI_XLSX_URL = "https://www.sipri.org/sites/default/files/SIPRI-Milex-data-1949-2025_v1.2.xlsx"

# IEA: new SDMX datasets. Exact dimension keys are intentionally supplied
# through the official mapping files rather than guessed here.
IEA_DATASETS = {
    "electricity_generation": "MESGEN",
    "electricity_balance": "MESBAL",
}

SIPRI_DATABASE_URL = "https://www.sipri.org/databases/milex"

# ---------------------------------------------------------------------------
# Aktualitätsregeln
# ---------------------------------------------------------------------------

# Diese Werte sind bewusst technisch-konservativ und beziehen sich auf die
# erwartete Veröffentlichungsfrequenz. Sie bedeuten NICHT, dass ein Datensatz
# nach Ablauf automatisch gelöscht wird. Bei Überschreitung bleibt er im Cache
# und wird nur als zu alt/unavailable behandelt, sofern kein neuer Abruf gelingt.
MAX_AGE_DAYS = {
    "OECD_STAN": 450,
    "OECD_PRODUCTIVITY": 450,
    "AI_INDEX": 500,
    "IEA_ELECTRICITY": 75,
    "SIPRI_DEFENCE": 500,
}

# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------

LOG = logging.getLogger("struktur_trends")


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None


def age_days(timestamp: str | None) -> float | None:
    dt = parse_iso(timestamp)
    if dt is None:
        return None
    return max(0.0, (datetime.now(timezone.utc) - dt).total_seconds() / 86400.0)


def http_get(url: str, timeout: int = 60, retries: int = 5) -> bytes:
    """Robuster öffentlicher HTTP-Abruf mit besonderer Behandlung von 429.

    OECD SDMX kann bei zu vielen Einzelabfragen temporär HTTP 429 liefern.
    Wir reagieren darauf mit Retry-After bzw. exponentiellem Backoff.
    Die eigentliche Entlastung erfolgt zusätzlich dadurch, dass OECD-Abfragen
    unten gebündelt werden und nicht mehr pro Land/Branche einzeln erfolgen.
    """
    last_error: Exception | None = None
    headers = {
        "User-Agent": "Neuber-Macro-Structural-Trends/1.1",
        "Accept": "text/csv,application/vnd.sdmx.data+csv;version=2.0.0,"
                  "application/csv,application/json;q=0.9,*/*;q=0.5",
    }

    for attempt in range(retries + 1):
        try:
            request = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            last_error = exc
            if exc.code != 429 or attempt >= retries:
                raise RuntimeError(
                    f"HTTP-Abruf fehlgeschlagen: {url} | HTTP {exc.code}: {exc.reason}"
                ) from exc

            retry_after = exc.headers.get("Retry-After")
            try:
                delay = float(retry_after) if retry_after else 0.0
            except ValueError:
                delay = 0.0

            # Wenn der Server keine Wartezeit vorgibt: konservativer Backoff.
            # Kleiner Jitter verhindert identische Retry-Zeitpunkte.
            if delay <= 0:
                delay = min(60.0, 5.0 * (2 ** attempt))
            delay += random.uniform(0.25, 1.25)
            LOG.warning(
                "HTTP 429 von OECD/Quelle; Retry %d/%d in %.1fs: %s",
                attempt + 1, retries, delay, url
            )
            time.sleep(delay)
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last_error = exc
            if attempt < retries:
                delay = min(30.0, 2.0 * (2 ** attempt)) + random.uniform(0.1, 0.8)
                time.sleep(delay)

    raise RuntimeError(f"HTTP-Abruf fehlgeschlagen: {url} | {last_error}")


def _download_temp(url: str, suffix: str = "") -> str:
    raw = http_get(url, timeout=180, retries=3)
    fd, path = tempfile.mkstemp(prefix="struktur_trends_", suffix=suffix)
    os.close(fd)
    Path(path).write_bytes(raw)
    return path


def _http_get_headers(url: str, headers: dict[str, str], timeout: int = 120, retries: int = 3) -> bytes:
    last_error = None
    base = {"User-Agent": "Neuber-Macro-Structural-Trends/1.2"}
    base.update(headers)
    for attempt in range(retries + 1):
        try:
            request = urllib.request.Request(url, headers=base)
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            last_error = exc
            if exc.code != 429 or attempt >= retries: raise RuntimeError(f"HTTP-Abruf fehlgeschlagen: {url} | HTTP {exc.code}: {exc.reason}") from exc
            retry_after = exc.headers.get("Retry-After")
            try: delay = float(retry_after) if retry_after else min(60.0, 5.0 * (2 ** attempt))
            except ValueError: delay = min(60.0, 5.0 * (2 ** attempt))
            time.sleep(delay + random.uniform(0.2, 1.0))
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            last_error = exc
            if attempt < retries: time.sleep(min(30.0, 2.0 * (2 ** attempt)) + random.uniform(0.1, 0.8))
    raise RuntimeError(f"HTTP-Abruf fehlgeschlagen: {url} | {last_error}")


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return new_cache()

    try:
        with path.open("r", encoding="utf-8") as fh:
            obj = json.load(fh)
        if not isinstance(obj, dict):
            raise ValueError("Cache ist kein JSON-Objekt.")
        return obj
    except (OSError, json.JSONDecodeError, ValueError) as exc:
        raise RuntimeError(f"Cache konnte nicht gelesen werden: {path} | {exc}") from exc


def atomic_write_json(path: Path, data: dict[str, Any]) -> None:
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8", newline="\n") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2, sort_keys=False)
        fh.write("\n")
        fh.flush()
    tmp.replace(path)


def empty_field(source: str, dataset: str, frequency: str) -> dict[str, Any]:
    return {
        "source": source,
        "dataset": dataset,
        "version": None,
        "frequency": frequency,
        "status": "UNAVAILABLE",
        "unit": None,
        "data_period": None,
        "retrieved_at": None,
        "last_successful_update": None,
        "series_count": 0,
        "observation_count": 0,
        "data": {},
        "source_notes": [],
    }


def new_cache() -> dict[str, Any]:
    return {
        "cache_version": CACHE_VERSION,
        "cache_created_at": now_iso(),
        "cache_updated_at": now_iso(),

        "OECD_STAN": {
            "value_added": empty_field("OECD", "DSD_STAN@DF_STAN_2025", "annual"),
            "investment": empty_field("OECD", "DSD_STAN@DF_STAN_2025", "annual"),
            "labour_input": empty_field("OECD", "DSD_STAN@DF_STAN_2025", "annual"),
        },

        "OECD_PRODUCTIVITY": {
            "labour_productivity_level": empty_field(
                "OECD", "DSD_PDB@DF_PDB", "annual"
            ),
            "labour_productivity_growth": empty_field(
                "OECD", "DSD_PDB@DF_PDB", "annual"
            ),
        },

        "AI_INDEX": {
            "ai_investment": empty_field("Stanford AI Index", "AI Index Public Data", "annual"),
            "ai_adoption": empty_field("Stanford AI Index", "AI Index Public Data", "annual"),
            "ai_compute": empty_field("Stanford AI Index", "AI Index Public Data", "annual"),
        },

        "IEA_ELECTRICITY": {
            "electricity_generation": empty_field("IEA", "MESGEN", "monthly"),
            "electricity_balance": empty_field("IEA", "MESBAL", "monthly"),
        },

        "SIPRI_DEFENCE": {
            "military_expenditure_real": empty_field(
                "SIPRI", "Military Expenditure Database", "annual"
            ),
            "military_burden_gdp": empty_field(
                "SIPRI", "Military Expenditure Database", "annual"
            ),
            "military_share_government": empty_field(
                "SIPRI", "Military Expenditure Database", "annual"
            ),
        },
    }


def merge_metadata(field: dict[str, Any], **updates: Any) -> dict[str, Any]:
    result = deepcopy(field)
    result.update(updates)
    return result


def set_real_cached_if_valid(field: dict[str, Any], source_key: str) -> dict[str, Any]:
    """
    Bei Abruffehler alten Datenbestand erhalten.
    Status nur dann REAL_CACHED, wenn ein erfolgreicher Abruf vorhanden
    und die definierte Altersgrenze eingehalten wird.
    """
    if not field.get("data"):
        result = deepcopy(field)
        result["status"] = "UNAVAILABLE"
        return result

    age = age_days(field.get("last_successful_update"))
    max_age = MAX_AGE_DAYS[source_key]

    result = deepcopy(field)
    if age is not None and age <= max_age:
        result["status"] = "REAL_CACHED"
    else:
        result["status"] = "UNAVAILABLE"
        result.setdefault("source_notes", []).append(
            f"Cache älter als zulässige {max_age} Tage."
        )
    return result


def cache_field_summary(field: dict[str, Any]) -> dict[str, Any]:
    return {
        "status": field.get("status"),
        "frequency": field.get("frequency"),
        "data_period": field.get("data_period"),
        "retrieved_at": field.get("retrieved_at"),
        "last_successful_update": field.get("last_successful_update"),
        "observations": sum(
            len(v) if isinstance(v, dict) else 1
            for v in field.get("data", {}).values()
        ),
    }


# ---------------------------------------------------------------------------
# OECD STAN
# ---------------------------------------------------------------------------

def oecd_stan_url(measure: str, start_period: int | None = None) -> str:
    """Bündelte STAN-Abfrage für alle benötigten Länder und Aktivitäten."""
    countries = "+".join(OECD_STAN_REFERENCE_AREAS)
    activities = "+".join(OECD_STAN_ACTIVITIES)
    key = f"A.{countries}.{activities}.{measure}.V.XDC"
    url = OECD_STAN_BASE + key + "?dimensionAtObservation=AllDimensions"
    if start_period:
        url += f"&startPeriod={start_period}"
    return url


def parse_csv_payload(raw: bytes) -> list[dict[str, str]]:
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [dict(row) for row in reader]


def first_present(row: dict[str, Any], *names: str) -> Any:
    normalized = {str(k).upper(): v for k, v in row.items()}
    for name in names:
        if name.upper() in normalized:
            return normalized[name.upper()]
    return None


def parse_sdmx_csv_rows(raw: bytes) -> list[dict[str, Any]]:
    """Toleranter Parser für SDMX-CSV mit erhaltenen Dimensionsmetadaten."""
    rows = parse_csv_payload(raw)
    output: list[dict[str, Any]] = []

    for row in rows:
        period = first_present(row, "TIME_PERIOD", "TIME")
        value = first_present(row, "OBS_VALUE", "VALUE")
        if period is None or value in (None, ""):
            continue

        try:
            numeric = float(str(value).replace(",", "."))
        except ValueError:
            continue

        output.append({
            "period": str(period),
            "value": numeric,
            "unit": first_present(row, "UNIT", "UNIT_MEASURE"),
            "reference_area": first_present(row, "REF_AREA", "REFERENCE_AREA"),
            "activity": first_present(row, "ACTIVITY", "ACTIVITY_CODE"),
            "measure": first_present(row, "MEASURE", "MEASURE_CODE"),
            "frequency": first_present(row, "FREQ", "FREQUENCY"),
            "transformation": first_present(row, "TRANSFORMATION"),
            "observation_status": first_present(
                row, "OBS_STATUS", "OBSERVATION_STATUS", "OBS_STATUS_CODE"
            ),
            "source_flag": first_present(
                row, "OBS_FLAG", "OBSERVATION_FLAG", "OBSERVATION_STATUS"
            ),
        })

    return output


def fetch_stan_measure(measure: str, start_period: int = 2000) -> list[dict[str, Any]]:
    url = oecd_stan_url(measure, start_period)
    raw = http_get(url)
    return parse_sdmx_csv_rows(raw)


def fetch_stan_labour_input(start_period: int = 2000) -> list[dict[str, Any]]:
    countries = "+".join(OECD_STAN_REFERENCE_AREAS)
    activities = "+".join(OECD_STAN_ACTIVITIES)
    # Official STAN explorer example uses A.FRA._T...H for Working Hours.
    key = f"A.{countries}.{activities}...H"
    url = OECD_STAN_BASE + key + "?dimensionAtObservation=AllDimensions"
    if start_period:
        url += f"&startPeriod={start_period}"
    raw = http_get(url)
    rows = parse_sdmx_csv_rows(raw)
    for row in rows:
        row["measure"] = "WORKING_HOURS"
    return rows


def _group_stan_rows(rows: list[dict[str, Any]], measure: str) -> dict[str, Any]:
    """Gruppiert STAN-Rohdaten und entfernt nur exakte Serien-Duplikate.

    Eine Serie wird durch Land + Aktivität definiert. Innerhalb der Serie
    bleiben unterschiedliche Perioden/Einheiten/Transformationen erhalten.
    Identische Beobachtungen werden deterministisch dedupliziert.
    """
    observations: dict[str, Any] = {}

    for row in rows:
        country = str(row.get("reference_area") or "").strip()
        activity = str(row.get("activity") or "").strip()

        if country not in OECD_STAN_REFERENCE_AREAS:
            continue
        if activity not in OECD_STAN_ACTIVITIES:
            continue

        key = f"{country}|{activity}"
        block = observations.setdefault(
            key,
            {
                "reference_area": country,
                "activity": activity,
                "activity_name": OECD_STAN_ACTIVITIES[activity],
                "measure": measure,
                "observations": [],
            },
        )

        observation = {
            k: v for k, v in row.items()
            if k not in {"reference_area", "activity", "measure"}
        }

        # Deduplizierung ohne Informationsverlust:
        # gleiche Periode + gleiche technische Metadaten + gleicher Wert
        # werden nur einmal gespeichert.
        dedupe_key = json.dumps(
            observation,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )

        seen = block.setdefault("_seen", set())
        if dedupe_key not in seen:
            seen.add(dedupe_key)
            block["observations"].append(observation)

    # Interne Dedupe-Hilfsmenge nicht in den Cache schreiben.
    for block in observations.values():
        block.pop("_seen", None)
        block["observations"].sort(
            key=lambda x: (
                str(x.get("period") or ""),
                str(x.get("unit") or ""),
                str(x.get("transformation") or ""),
            )
        )

    return observations


def _count_observations(container: dict[str, Any]) -> int:
    return sum(
        len(block.get("observations", []))
        for block in container.values()
    )


def _audit_stan_container(container: dict[str, Any], measure: str) -> tuple[int, int]:
    """Prüft, dass nur konfigurierte STAN-Serien im Cache landen."""
    series_count = 0
    observation_count = 0

    for key, block in container.items():
        country = block.get("reference_area")
        activity = block.get("activity")

        if country not in OECD_STAN_REFERENCE_AREAS:
            raise RuntimeError(f"Unerwartetes STAN-Land im Cache: {country!r}")
        if activity not in OECD_STAN_ACTIVITIES:
            raise RuntimeError(f"Unerwartete STAN-Aktivität im Cache: {activity!r}")
        if block.get("measure") != measure:
            raise RuntimeError(
                f"Falsches STAN-Measure in Serie {key}: "
                f"{block.get('measure')!r} statt {measure!r}"
            )

        series_count += 1
        observation_count += len(block.get("observations", []))

    if series_count == 0 or observation_count == 0:
        raise RuntimeError(
            f"STAN-Audit fehlgeschlagen: measure={measure}, "
            f"series={series_count}, observations={observation_count}"
        )

    return series_count, observation_count


def update_oecd_stan(cache: dict[str, Any], start_period: int = 2000) -> None:
    mapping = {
        "value_added": "B1G",
        "investment": "P51G",
    }

    for field_name, measure in mapping.items():
        field = cache["OECD_STAN"][field_name]
        retrieved = now_iso()
        try:
            rows = fetch_stan_measure(measure, start_period=start_period)
            observations = _group_stan_rows(rows, measure)
            series_count, observation_count = _audit_stan_container(
                observations, measure
            )

            field.update({
                "version": "1.0",
                "status": "REAL",
                "retrieved_at": retrieved,
                "last_successful_update": retrieved,
                "data_period": _latest_period(observations),
                "data": observations,
                "series_count": series_count,
                "observation_count": observation_count,
                "source_notes": [
                    "STAN-Abfrage gebündelt über konfigurierte Länder und Aktivitäten.",
                    "Exakte Beobachtungsduplikate werden vor dem Cache entfernt.",
                ],
            })

        except Exception as exc:
            LOG.warning("OECD STAN %s: %s", field_name, exc)
            cache["OECD_STAN"][field_name] = set_real_cached_if_valid(
                field, "OECD_STAN"
            )

    labour = cache["OECD_STAN"]["labour_input"]
    try:
        rows = fetch_stan_labour_input(start_period=start_period)
        grouped = _group_stan_rows(rows, "WORKING_HOURS")
        if grouped:
            ts = now_iso()
            labour.update({
                "version": "1.0",
                "status": "REAL",
                "retrieved_at": ts,
                "last_successful_update": ts,
                "data_period": _latest_period(grouped),
                "data": grouped,
                "series_count": len(grouped),
                "observation_count": _count_observations(grouped),
                "source_notes": [
                    "STAN labour input über die offizielle Working-Hours-Einheit H selektiert; kein gerateter Measure-Code.",
                ],
            })
        else:
            raise RuntimeError("OECD STAN lieferte keine Working-Hours-Beobachtungen.")
    except Exception as exc:
        LOG.warning("OECD STAN labour_input: %s", exc)
        cache["OECD_STAN"]["labour_input"] = set_real_cached_if_valid(
            labour, "OECD_STAN"
        )

def _latest_period(container: dict[str, Any]) -> str | None:
    periods: list[str] = []
    for block in container.values():
        for row in block.get("observations", []):
            p = row.get("period")
            if p:
                periods.append(str(p))
    return max(periods) if periods else None


# ---------------------------------------------------------------------------
# OECD Productivity
# ---------------------------------------------------------------------------

def update_oecd_productivity(cache: dict[str, Any], start_period: int = 2000) -> None:
    """Produktivitäts-Level als eine gebündelte OECD-Abfrage."""
    countries = "+".join(OECD_PRODUCTIVITY_REFERENCE_AREAS)
    key = f"{countries}.A.{OECD_PRODUCTIVITY_MEASURE}._T.XDC_H..N.."
    url = (
        OECD_PRODUCTIVITY_BASE + key
        + "?dimensionAtObservation=AllDimensions"
        + f"&startPeriod={start_period}"
    )

    level = cache["OECD_PRODUCTIVITY"]["labour_productivity_level"]
    try:
        raw = http_get(url)
        rows = parse_sdmx_csv_rows(raw)
        grouped: dict[str, Any] = {}
        for row in rows:
            country = str(row.get("reference_area") or "").strip()
            if country not in OECD_PRODUCTIVITY_REFERENCE_AREAS:
                continue
            grouped.setdefault(
                country,
                {
                    "reference_area": country,
                    "measure": OECD_PRODUCTIVITY_MEASURE,
                    "observations": [],
                },
            )["observations"].append({
                k: v for k, v in row.items() if k != "reference_area"
            })

        if grouped:
            ts = now_iso()
            observation_count = _count_observations(grouped)
            level.update({
                "version": "2.0",
                "status": "REAL",
                "retrieved_at": ts,
                "last_successful_update": ts,
                "data_period": _latest_period(grouped),
                "series_count": len(grouped),
                "observation_count": observation_count,
                "data": grouped,
                "source_notes": [
                    "Gebündelte OECD-Produktivitätsabfrage über die konfigurierte Länderliste.",
                ],
            })
        else:
            raise RuntimeError("OECD Productivity lieferte keine verwertbaren Beobachtungen.")
    except Exception as exc:
        LOG.warning("OECD Productivity: %s", exc)
        cache["OECD_PRODUCTIVITY"]["labour_productivity_level"] = (
            set_real_cached_if_valid(level, "OECD_PRODUCTIVITY")
        )

    growth = cache["OECD_PRODUCTIVITY"]["labour_productivity_growth"]
    try:
        countries = "+".join(OECD_PRODUCTIVITY_REFERENCE_AREAS)
        # Official OECD growth-rate database. GDP per hour worked is the
        # published labour-productivity growth measure; Python does not
        # calculate the growth itself.
        key = f"{countries}.A.GDPHRS..PA..GY.."
        url = OECD_PRODUCTIVITY_GROWTH_BASE + key
        url += "?dimensionAtObservation=AllDimensions"
        if start_period:
            url += f"&startPeriod={start_period}"
        raw = http_get(url)
        rows = parse_sdmx_csv_rows(raw)
        grouped = {}
        for row in rows:
            country = str(row.get("reference_area") or "").strip()
            if country not in OECD_PRODUCTIVITY_REFERENCE_AREAS:
                continue
            grouped.setdefault(country, {
                "reference_area": country,
                "measure": "GDPHRS",
                "observations": [],
            })["observations"].append({
                k: v for k, v in row.items() if k != "reference_area"
            })
        if not grouped:
            raise RuntimeError("OECD Productivity Growth lieferte keine Beobachtungen.")
        ts = now_iso()
        growth.update({
            "version": "1.0",
            "status": "REAL",
            "retrieved_at": ts,
            "last_successful_update": ts,
            "data_period": _latest_period(grouped),
            "data": grouped,
            "series_count": len(grouped),
            "observation_count": _count_observations(grouped),
            "source_notes": [
                "Offizielle OECD Productivity Growth Rates: GDP per hour worked, Growth rate over 1 year.",
                "Wachstum wird nicht in Python berechnet.",
            ],
        })
    except Exception as exc:
        LOG.warning("OECD Productivity Growth: %s", exc)
        cache["OECD_PRODUCTIVITY"]["labour_productivity_growth"] = (
            set_real_cached_if_valid(growth, "OECD_PRODUCTIVITY")
        )


# ---------------------------------------------------------------------------
# Stanford AI Index
# ---------------------------------------------------------------------------

def _fetch_text(url: str) -> str:
    return http_get(url, timeout=60).decode("utf-8", errors="replace")


def _extract_first_number(text: str, patterns: list[str]) -> float | None:
    for pattern in patterns:
        m = re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL)
        if m:
            raw = m.group(1).replace(",", "")
            try:
                return float(raw)
            except ValueError:
                pass
    return None


def update_ai_index(cache: dict[str, Any]) -> None:
    """Liest aktuelle, ausdrücklich veröffentlichte Stanford-AI-Index-Werte.

    Es werden keine Werte im Code fest verdrahtet. Die offiziellen Stanford-
    Seiten werden bei jedem Lauf abgefragt; fällt der Abruf aus, greift die
    normale REAL_CACHED/UNAVAILABLE-Logik.
    """
    ts = now_iso()
    try:
        economy = _fetch_text(STANFORD_ECONOMY_URL)
        rd = _fetch_text(STANFORD_RD_URL)
        takeaway = _fetch_text(STANFORD_TAKEAWAY_URL)

        investment = _extract_first_number(
            takeaway + "\n" + economy,
            [r"global corporate AI investments? hit \$([0-9.]+)\s*billion",
             r"global corporate AI investment[^$]{0,300}\$([0-9.]+)\s*billion",
             r"AI investment[^$]{0,300}\$([0-9.]+)\s*billion"],
        )
        adoption = _extract_first_number(
            economy,
            [r"organizational adoption.*?([0-9.]+)%",
             r"AI adoption.*?([0-9.]+)%"],
        )
        compute = _extract_first_number(
            rd,
            [r"reaching ([0-9.]+) million H100-equivalents",
             r"([0-9.]+) million H100-equivalents"],
        )

        results = {
            "ai_investment": (investment, "USD billion", "2025"),
            "ai_adoption": (adoption, "percent of surveyed organizations", "2025"),
            "ai_compute": (compute, "million H100-equivalents", "2025"),
        }
        for field_name, (value, unit, period) in results.items():
            field = cache["AI_INDEX"][field_name]
            if value is None:
                raise RuntimeError(f"Stanford-Wert nicht eindeutig extrahierbar: {field_name}")
            field.update({
                "version": "2026",
                "status": "REAL",
                "unit": unit,
                "data_period": period,
                "retrieved_at": ts,
                "last_successful_update": ts,
                "series_count": 1,
                "observation_count": 1,
                "data": {period: {"value": value, "source": "Stanford HAI AI Index 2026"}},
                "source_notes": [
                    "Offizielle Stanford HAI AI Index 2026-Webquelle; Wert wird zur Laufzeit extrahiert.",
                    "Kein Python-Seitentrend oder Interpretationssignal.",
                ],
            })
    except Exception as exc:
        LOG.warning("Stanford AI Index: %s", exc)
        for field_name in ("ai_investment", "ai_adoption", "ai_compute"):
            field = cache["AI_INDEX"][field_name]
            cache["AI_INDEX"][field_name] = set_real_cached_if_valid(field, "AI_INDEX")


# ---------------------------------------------------------------------------
# IEA MESGEN / MESBAL
# ---------------------------------------------------------------------------


IEA_MES_PAGE = "https://www.iea.org/data-and-statistics/data-product/monthly-electricity-statistics"
IEA_MES_DOCUMENTATION_PDF = "https://iea.blob.core.windows.net/assets/a7d1b044-38cd-4b85-a804-c68be5b45687/Monthly_electricity_statistics_Documentation_2026.pdf"
IEA_MES_MAPPING_SUMMARY_PDF = "https://iea.blob.core.windows.net/assets/2ae5c8ac-0397-4e85-ac5d-5ec9be3e7c01/MESSDMXmappingsummary2026.pdf"


def _iea_parse_csv_bytes(raw: bytes) -> list[dict[str, Any]]:
    """Parse official IEA SDMX CSV/ZIP content.

    A release ZIP can contain more than one CSV (for example data plus
    metadata).  Score all CSV candidates and select the file that actually
    looks like an SDMX observation table instead of blindly taking the first
    CSV.
    """
    import csv, io, zipfile

    payloads: list[tuple[str, bytes]] = []
    if raw[:2] == b"PK":
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for name in zf.namelist():
                if name.lower().endswith((".csv", ".txt")):
                    payloads.append((name, zf.read(name)))
    else:
        payloads.append(("payload", raw))

    best_rows: list[dict[str, Any]] = []
    best_score = -1
    best_name = ""
    for name, payload in payloads:
        text = payload.decode("utf-8-sig", errors="replace")
        if not text.strip():
            continue
        try:
            dialect = csv.Sniffer().sniff(text[:16384], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        try:
            rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        except Exception:
            continue
        if not rows:
            continue
        normalized = [{str(k).strip(): v for k, v in row.items()} for row in rows]
        headers = {str(k).strip().upper() for k in normalized[0].keys()}
        score = 0
        for required in ("COUNTRY", "TIME_PERIOD", "OBS_VALUE"):
            if required in headers:
                score += 20
        if "ENERGY_BALANCE_FLOW" in headers:
            score += 10
        if "ENERGY_PRODUCT" in headers:
            score += 10
        if "FREQUENCY" in headers:
            score += 5
        if "UNIT" in headers:
            score += 5
        if "CONF_STATUS" in headers or "QUALIFIER" in headers:
            score += 3
        score += min(len(normalized) / 1000.0, 10.0)
        if score > best_score:
            best_score = score
            best_rows = normalized
            best_name = name

    if best_rows:
        LOG.info("IEA SDMX CSV ausgewählt: %s | rows=%s | score=%.1f", best_name, len(best_rows), best_score)
    return best_rows

def _iea_parse_json_bytes(raw: bytes) -> list[dict[str, Any]]:
    try:
        obj = json.loads(raw.decode("utf-8-sig", errors="replace"))
    except Exception:
        return []
    if isinstance(obj, list) and obj and all(isinstance(x, dict) for x in obj):
        return [{str(k): v for k, v in x.items()} for x in obj]
    if not isinstance(obj, dict):
        return []
    structure = obj.get("structure", {}) if isinstance(obj.get("structure"), dict) else {}
    dims = structure.get("dimensions", {}) if isinstance(structure.get("dimensions"), dict) else {}
    series_dims = dims.get("series", []) or []
    obs_dims = dims.get("observation", []) or []
    datasets = obj.get("dataSets") or obj.get("data") or []
    if not datasets or not isinstance(datasets[0], dict):
        return []
    ds = datasets[0]
    def dim_value(dim, index):
        values = dim.get("values", []) if isinstance(dim, dict) else []
        if index >= len(values):
            return index
        item = values[index]
        return item.get("id", item.get("name", index)) if isinstance(item, dict) else item
    out = []
    series = ds.get("series", {})
    if isinstance(series, dict):
        for series_key, series_obj in series.items():
            if not isinstance(series_obj, dict):
                continue
            sk = [int(x) for x in str(series_key).split(":") if str(x).isdigit()]
            base = {}
            for i, dim in enumerate(series_dims):
                base[dim.get("id", f"SERIES_{i}")] = dim_value(dim, sk[i] if i < len(sk) else 0)
            observations = series_obj.get("observations", {})
            if isinstance(observations, dict):
                for obs_key, obs in observations.items():
                    oi = [int(x) for x in str(obs_key).split(":") if str(x).isdigit()]
                    row = dict(base)
                    for i, dim in enumerate(obs_dims):
                        row[dim.get("id", f"OBS_{i}")] = dim_value(dim, oi[i] if i < len(oi) else 0)
                    row["OBS_VALUE"] = obs[0] if isinstance(obs, list) and obs else obs
                    if isinstance(obs, list) and len(obs) > 1:
                        row["OBS_STATUS"] = obs[1]
                    out.append(row)
    return out


def _iea_parse_xml_bytes(raw: bytes) -> list[dict[str, Any]]:
    import xml.etree.ElementTree as ET
    try:
        root = ET.fromstring(raw)
    except Exception:
        return []
    rows = []
    for series in root.iter():
        if not series.tag.lower().endswith("series"):
            continue
        base = {k.split("}")[-1]: v for k, v in series.attrib.items()}
        for child in series:
            if child.tag.lower().endswith("obs"):
                row = dict(base)
                row.update({k.split("}")[-1]: v for k, v in child.attrib.items()})
                rows.append(row)
    return rows


def _iea_parse_payload_variants(raw: bytes) -> list[tuple[str, list[dict[str, Any]]]]:
    """Parse one payload several ways; the payload itself is never downloaded twice."""
    parsers = (
        ("csv", _iea_parse_csv_bytes),
        ("json", _iea_parse_json_bytes),
        ("xml", _iea_parse_xml_bytes),
    )
    result = []
    for name, parser in parsers:
        try:
            rows = parser(raw)
            if rows:
                result.append((name, rows))
        except Exception as exc:
            LOG.debug("IEA Parser %s fehlgeschlagen: %s", name, exc)
    return result


def _iea_html_links(raw: bytes, flow: str) -> list[str]:
    """Extract official download links from the IEA MES page.

    The page is the authoritative source for the current release.  We do not
    invent asset IDs or filenames.  The parser accepts normal hrefs as well as
    JSON/escaped links emitted by the page frontend.
    """
    text = raw.decode("utf-8", errors="replace")
    text = text.replace("\\u0026", "&").replace("\\/", "/").replace("&amp;", "&")
    candidates = []
    patterns = (
        r'https?://[^\"\'<>\s]+',
        r'href=[\"\']([^\"\']+)[\"\']',
        r'(?:url|href|downloadUrl|download_url|fileUrl|file_url)[\"\']?\s*[:=]\s*[\"\']([^\"\']+)[\"\']',
    )
    for pattern in patterns:
        candidates.extend(re.findall(pattern, text, flags=re.I))
    result = []
    flow_upper = flow.upper()
    for url in candidates:
        if not isinstance(url, str):
            continue
        url = url.replace("\\u002F", "/").replace("&amp;", "&")
        if url.startswith("/"):
            url = "https://www.iea.org" + url
        low = url.lower()
        if flow_upper not in url.upper() and not any(token in low for token in ("monthly_electricity", "mesgen", "mesbal", "monthly-electricity-statistics")):
            continue
        if not any(token in low for token in (".zip", ".csv", ".txt", "download", "asset")):
            continue
        result.append(url)
    return list(dict.fromkeys(result))


def _iea_page_fetch_variants() -> list[tuple[str, bytes]]:
    """Fetch the official MES page with browser-like variants.

    GitHub runners have historically received HTTP 403 from the IEA website.
    We therefore try several ordinary browser headers before giving up.  No
    third-party mirror is used.
    """
    variants = (
        {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36", "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.9", "Referer": "https://www.google.com/"},
        {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 Version/18.6 Safari/605.1.15", "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8", "Accept-Language": "en-US,en;q=0.8"},
        {"User-Agent": "curl/8.10.1", "Accept": "text/html,*/*;q=0.8"},
        {"User-Agent": "python-urllib/3.13", "Accept": "text/html,*/*;q=0.5"},
    )
    out = []
    for i, headers in enumerate(variants, start=1):
        try:
            raw = _http_get_headers(IEA_MES_PAGE, headers, timeout=60, retries=1)
            if raw:
                out.append((f"page-header-{i}", raw))
                LOG.info("IEA MES-Seite erreichbar mit Header-Variante %s", i)
                break
        except Exception as exc:
            LOG.info("IEA MES-Seite Variante %s nicht verfügbar: %s", i, exc)
    return out


def _iea_mapping_metadata(flow: str) -> dict[str, Any]:
    """Analyse the complete official IEA mapping workbook.

    This Lauf is deliberately a structure-analysis step.  We do not invent a
    data URL and we do not turn mapping metadata into observations.  Instead,
    the complete workbook is inspected so that the next implementation can
    derive the real MESGEN/MESBAL structure from the IEA's own mapping.
    """
    url = IEA_MAPPING_URLS[flow]
    try:
        import openpyxl
        raw = _http_get_headers(
            url,
            {"Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*;q=0.2"},
            timeout=60,
            retries=2,
        )
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        result: dict[str, Any] = {"url": url, "sheets": wb.sheetnames, "sheet_info": {}}

        LOG.info("IEA %s Mapping geladen: %s", flow, ", ".join(wb.sheetnames))

        url_re = re.compile(r"https?://[^\s<>\"']+", re.I)
        keyword_re = re.compile(
            r"dataflow|datastructure|data structure|dimension|attribute|series|observation|"
            r"frequency|time period|time_period|ref_area|reference area|measure|unit|"
            r"generation|balance|download|\.zip|\.csv|sdmx|estat|stat",
            re.I,
        )
        code_re = re.compile(r"^[A-Za-z][A-Za-z0-9_.-]{1,31}$")

        for ws in wb.worksheets:
            # Read the complete used range, not just the first 12 rows.  Mapping
            # workbooks often place the actual dimension/code information well
            # below the explanatory header.
            rows: list[list[Any]] = []
            max_cols = 0
            nonempty = 0
            url_candidates: list[str] = []
            keyword_hits: list[tuple[int, int, str]] = []
            code_samples: list[str] = []

            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = list(row)
                max_cols = max(max_cols, len(values))
                cleaned = [str(v).strip() if v is not None else "" for v in values]
                if any(cleaned):
                    nonempty += 1
                rows.append(cleaned)
                for c_idx, value in enumerate(cleaned, start=1):
                    if not value:
                        continue
                    for found in url_re.findall(value):
                        url_candidates.append(found.rstrip(').,;'))
                    if keyword_re.search(value):
                        keyword_hits.append((r_idx, c_idx, value[:180]))
                    if len(code_samples) < 80 and code_re.fullmatch(value) and value.upper() not in {"CODE", "VALUE", "NAME", "LABEL", "ID"}:
                        code_samples.append(value)

            # Compact structural summary for the GitHub log.
            header_candidates = []
            for r_idx, row in enumerate(rows[:80], start=1):
                filled = [v for v in row if v]
                if len(filled) >= 2:
                    header_candidates.append({"row": r_idx, "values": filled[:16]})

            info = {
                "rows": len(rows),
                "nonempty_rows": nonempty,
                "columns": max_cols,
                "url_candidates": list(dict.fromkeys(url_candidates))[:30],
                "keyword_hits": keyword_hits[:80],
                "code_samples": list(dict.fromkeys(code_samples))[:80],
                "header_candidates": header_candidates[:12],
            }
            result["sheet_info"][ws.title] = info

            LOG.info(
                "IEA %s Mapping Sheet '%s': rows=%s nonempty=%s cols=%s",
                flow, ws.title, len(rows), nonempty, max_cols,
            )
            if info["url_candidates"]:
                LOG.info("IEA %s %s URLs: %s", flow, ws.title, " | ".join(info["url_candidates"][:10]))
            if info["header_candidates"]:
                for header in info["header_candidates"][:4]:
                    LOG.info("IEA %s %s Header-Kandidat Zeile %s: %s", flow, ws.title, header["row"], header["values"])
            if info["keyword_hits"]:
                for hit in info["keyword_hits"][:12]:
                    LOG.info("IEA %s %s Struktur-Treffer r=%s c=%s: %s", flow, ws.title, hit[0], hit[1], hit[2])
            if info["code_samples"]:
                LOG.info("IEA %s %s Code-Beispiele: %s", flow, ws.title, ", ".join(info["code_samples"][:30]))

        return result
    except Exception as exc:
        LOG.info("IEA %s Mapping-Analyse nicht verfügbar: %s", flow, exc)
        return {}


def _group_iea_rows(rows: list[dict[str, Any]], dataset: str) -> dict[str, Any]:
    grouped: dict[str, Any] = {}
    for row in rows:
        lowered = {str(k).strip().lower(): v for k, v in row.items()}
        def pick(*names):
            for name in names:
                if name in lowered and str(lowered[name] or "").strip():
                    return str(lowered[name]).strip()
            return None
        country = pick("country", "ref_area", "reference area", "reference_area", "country or area", "economy", "geo", "area")
        period = pick("time_period", "time period", "time_period_start", "period", "time", "date")
        if not country or not period:
            continue
        grouped.setdefault(country, {"reference_area": country, "dataset": dataset, "observations": []})["observations"].append(dict(row))
    return grouped

def _iea_stable_structure_candidates(flow: str) -> list[tuple[str, str]]:
    """Discover the real IEA .Stat dataflow instead of assuming MESGEN/MESBAL IDs.

    The IEA stable host exposes a standard .Stat/SDMX registry.  MESGEN/MESBAL
    are the published file labels/mapping names, not proven dataflow IDs.  The
    first candidates therefore enumerate all dataflows and let
    ``_iea_extract_resource_ids`` identify matching flows by title/ID.
    Specific legacy guesses remain only as a narrow fallback.
    """
    flow = str(flow).upper()
    out = []
    bases = (
        (IEA_API_BASE_STABLE, "stable-api"),
        (IEA_API_BASE_NSI_STABLE, "stable-nsi"),
    )
    for base, host_label in bases:
        # .Stat standard structural query: all agencies, all dataflows, latest.
        out.append((f"{host_label}-all", f"{base}/dataflow/all/all/latest?detail=allstubs"))
        out.append((f"{host_label}-all-v", f"{base}/dataflow/all/all/all?detail=allstubs"))
        # Some .Stat deployments expose the compact all-dataflow form.
        out.append((f"{host_label}-all-short", f"{base}/dataflow/all/latest"))
        # Narrow fallbacks only; these are not treated as authoritative IDs.
        for agency in ("OECD.IEA", "IEA"):
            for version in ("1.0", "latest", "2026"):
                out.append((f"{host_label}-guess-v1", f"{base}/dataflow/{agency}/{flow}/{version}"))
    return list(dict.fromkeys(out))


def _iea_find_dataflow(flow: str) -> list[tuple[str, str, bytes]]:
    """Find the real MES resource on the official IEA SDMX service.

    A data response is also a valid discovery signal for this public service,
    so this function accepts both structure and dataflow-style responses.
    """
    found = []
    seen = set()
    headers = {
        "Accept": "application/vnd.sdmx.structure+csv;version=2.0.0,application/vnd.sdmx.structure+xml,application/vnd.sdmx.structure+json,application/vnd.sdmx.data+csv;version=2.0.0,application/json,application/xml,text/csv;q=0.8,*/*;q=0.2",
        "User-Agent": "Mozilla/5.0 (compatible; StrukturTrends/1.0; +https://github.com/)"
    }
    for label, url in _iea_stable_structure_candidates(flow):
        try:
            raw = _http_get_headers(url, headers, timeout=90, retries=2)
            head = raw[:500].lower()
            if not raw or head.startswith(b"<!doctype html") or b"<html" in head:
                LOG.info("IEA %s Discovery %s: unerwartete HTML-Antwort | %s", flow, label, url)
                continue
            digest = hashlib.sha256(raw).hexdigest()
            if digest in seen:
                continue
            seen.add(digest)
            found.append((label, url, raw))
            LOG.info("IEA %s Discovery erreichbar: %s | bytes=%s | head=%r", flow, url, len(raw), raw[:120])
        except Exception as exc:
            LOG.info("IEA %s Discovery %s nicht verfügbar: %s", flow, label, exc)
    return found


def _iea_data_candidates(flow: str, structures: list[tuple[str, str, bytes]]) -> list[tuple[str, str]]:
    """Create official SDMX data candidates without inventing dimension codes.

    The key point is that SDMX v1 accepts ``all`` as the key wildcard and v2
    permits an omitted key.  The previous implementation used strings of dots,
    which produced HTTP 400 on the IEA service.  We therefore test the standards-
    compliant forms first and use component filters in v2 where supported.
    """
    triples = []
    for _, _, raw in structures:
        triples.extend(_iea_extract_resource_ids(raw, flow))
    # Always test the documented IEA mapping agency and the public IEA agency,
    # with explicit versions first.  The mapping workbook identifies its
    # codelists as OECD.IEA, so that agency is the primary candidate.
    for agency in ("OECD.IEA", "IEA"):
        for version in ("1.0", "latest", "2026"):
            triples.append((agency, flow, version))
    triples = list(dict.fromkeys(triples))

    urls = []
    for agency, resource, version in triples:
        # SDMX REST v1: flowRef/key/providerRef.  ``all`` is the standards-
        # defined key wildcard; providerRef is omitted first.
        for base, label in ((IEA_API_BASE_STABLE, "stable-v1"), (IEA_API_BASE_NSI_STABLE, "stable-nsi-v1"), (IEA_API_BASE, "legacy-v1")):
            for key in ("all", ""):
                key_part = f"/{key}" if key else ""
                for provider in ("", "/all", "/IEA", "/OECD.IEA"):
                    if provider and key == "" and provider == "/all":
                        pass
                    url = f"{base}/data/{agency},{resource},{version}{key_part}{provider}"
                    urls.append((url, f"{label}:{agency}/{resource}/{version}|key={key or 'omitted'}|provider={provider.strip('/') or 'omitted'}"))
                    urls.append((url + "?startPeriod=2020-01&endPeriod=2026-12&dimensionAtObservation=AllDimensions", f"{label}:{agency}/{resource}/{version}|key={key or 'omitted'}|provider={provider.strip('/') or 'omitted'}|period"))
        # SDMX REST v2: data/dataflow/{agency}/{resource}/{version}/{key};
        # omitted key means the whole dataflow. Component filters are preferred.
        base2 = f"{IEA_API_BASE_STABLE}/v2/data/dataflow/{agency}/{resource}/{version}"
        base2_nsi = f"{IEA_API_BASE_NSI_STABLE}/v2/data/dataflow/{agency}/{resource}/{version}"
        urls.append((base2, f"stable-v2:{agency}/{resource}/{version}|key=omitted"))
        urls.append((base2_nsi, f"stable-nsi-v2:{agency}/{resource}/{version}|key=omitted"))
        urls.append((base2 + "?startPeriod=2020-01&endPeriod=2026-12&c%5BFREQUENCY%5D=M", f"stable-v2:{agency}/{resource}/{version}|key=omitted|FREQUENCY=M"))
        urls.append((base2_nsi + "?startPeriod=2020-01&endPeriod=2026-12&c%5BFREQUENCY%5D=M", f"stable-nsi-v2:{agency}/{resource}/{version}|key=omitted|FREQUENCY=M"))
        urls.append((base2 + "/all?startPeriod=2020-01&endPeriod=2026-12", f"stable-v2:{agency}/{resource}/{version}|key=all|period"))
        urls.append((base2_nsi + "/all?startPeriod=2020-01&endPeriod=2026-12", f"stable-nsi-v2:{agency}/{resource}/{version}|key=all|period"))
    return list(dict.fromkeys(urls))

def _iea_extract_rows(flow: str) -> tuple[list[dict[str, Any]], str]:
    """Acquire MESGEN/MESBAL through the official stable IEA SDMX service first.

    Only IEA-owned endpoints are used.  The older IEA page/ZIP route remains a
    fallback because the page is currently blocked by some GitHub runners.
    """
    errors: list[str] = []
    payload_candidates: list[tuple[str, bytes]] = []

    # 1) Official stable SDMX registry -> real flow -> data.
    structures = _iea_find_dataflow(flow)
    for structure_label, structure_url, structure_raw in structures:
        LOG.info("IEA %s Strukturquelle: %s", flow, structure_url)
    for url, identity in _iea_data_candidates(flow, structures):
        try:
            raw = _http_get_headers(url, {
                "Accept": "text/csv,application/vnd.sdmx.data+csv;version=2.0.0,application/vnd.sdmx.data+json,application/json,application/xml,application/zip;q=0.9,*/*;q=0.2",
                "User-Agent": "Mozilla/5.0 (compatible; StrukturTrends/1.0; +https://github.com/)"
            }, timeout=240, retries=2)
            head = raw[:300].lower()
            if raw and b"<html" not in head and not head.startswith(b"<!doctype"):
                payload_candidates.append((f"stable-sdmx:{identity}:{url}", raw))
                LOG.info("IEA %s SDMX-Datenpayload geladen: %s | bytes=%s | ZIP=%s", flow, url, len(raw), raw[:2] == b"PK")
            else:
                errors.append(f"stable-data {url}: HTML/leer")
        except Exception as exc:
            errors.append(f"stable-data {url}: {exc}")

    # 2) Official mappings remain a validation/structure source, never observations.
    mapping = _iea_mapping_metadata(flow)
    if mapping:
        LOG.info("IEA %s Mapping-Analyse vorhanden: %s", flow, ", ".join(mapping.get("sheets", [])))

    # 3) Fallback: official MES page and its current download links.
    if not payload_candidates:
        page_variants = _iea_page_fetch_variants()
        for page_label, page_raw in page_variants:
            urls = _iea_html_links(page_raw, flow)
            LOG.info("IEA %s: %s offizielle Download-Link-Kandidaten gefunden", flow, len(urls))
            for url in urls:
                try:
                    raw = _http_get_headers(url, {
                        "Accept": "application/zip,application/octet-stream,text/csv,*/*;q=0.2",
                        "Referer": IEA_MES_PAGE,
                    }, timeout=240, retries=2)
                    head = raw[:300].lower()
                    if raw[:2] == b"PK" or b"COUNTRY" in raw[:4096].upper():
                        payload_candidates.append((f"official-page-link:{page_label}:{url}", raw))
                        LOG.info("IEA %s offizielles Datenpayload geladen: %s | bytes=%s | ZIP=%s", flow, url, len(raw), raw[:2] == b"PK")
                    else:
                        LOG.warning("IEA %s Link lieferte unerwarteten Inhalt: %s | bytes=%s | head=%r", flow, url, len(raw), raw[:120])
                except Exception as exc:
                    errors.append(f"official-link {url}: {exc}")

    if not payload_candidates:
        raise RuntimeError(
            "IEA %s: offizielles SDMX-Datenpayload nicht erreichbar. "
            "Stable-Registry/Data und offizieller MES-Fallback erfolglos. %s"
            % (flow, " | ".join(errors[:20]))
        )

    unique_payloads = []
    seen = set()
    for label, raw in payload_candidates:
        digest = hashlib.sha256(raw).hexdigest()
        if digest not in seen:
            seen.add(digest)
            unique_payloads.append((label, raw))

    parser_orders = (
        ("csv", "json", "xml"), ("json", "csv", "xml"), ("xml", "csv", "json"),
        ("csv", "xml", "json"), ("json", "xml", "csv"),
    )
    parser_map = {"csv": _iea_parse_csv_bytes, "json": _iea_parse_json_bytes, "xml": _iea_parse_xml_bytes}
    strategy = 0
    for payload_index, (label, raw) in enumerate(unique_payloads[:4], start=1):
        for order in parser_orders:
            strategy += 1
            try:
                rows, parser_used = [], ""
                for name in order:
                    rows = parser_map[name](raw)
                    if rows:
                        parser_used = name
                        break
                grouped = _group_iea_rows(rows, flow)
                obs = sum(len(v.get("observations", [])) for v in grouped.values())
                if len(grouped) >= 2 and obs >= 10:
                    LOG.info("IEA %s erfolgreiche Strategie %s: %s | parser=%s | series=%s | obs=%s", flow, strategy, label, parser_used, len(grouped), obs)
                    return rows, f"strategy={strategy}:{label}:parser={parser_used}"
                errors.append(f"strategy {strategy}: unzureichende Struktur series={len(grouped)} obs={obs}")
            except Exception as exc:
                errors.append(f"strategy {strategy}: {exc}")

    raise RuntimeError("IEA %s: Parserstrategien erfolglos; %s" % (flow, " | ".join(errors[:20])))

def update_iea(cache: dict[str, Any]) -> None:
    for field_name, dataset in IEA_DATASETS.items():
        field = cache["IEA_ELECTRICITY"][field_name]
        try:
            rows, source = _iea_extract_rows(dataset)
            grouped = _group_iea_rows(rows, dataset)
            if not grouped:
                raise RuntimeError("IEA: Parser lieferte keine verwertbaren Zeitreihen")
            ts = now_iso()
            field.update({"version": "2026", "status": "REAL", "retrieved_at": ts, "last_successful_update": ts, "data_period": _latest_period(grouped), "data": grouped, "series_count": len(grouped), "observation_count": _count_observations(grouped), "source_notes": [f"Offizielle IEA Monthly Electricity Statistics, {dataset}.", f"Erfolgreiche Abruf-/Parserstrategie: {source}.", "IEA-Daten können quellenseitige Schätz-/Imputationsflags enthalten; diese werden erhalten."]})
        except Exception as exc:
            LOG.warning("IEA %s: %s", dataset, exc)
            cache["IEA_ELECTRICITY"][field_name] = set_real_cached_if_valid(field, "IEA_ELECTRICITY")


# ---------------------------------------------------------------------------
# SIPRI
# ---------------------------------------------------------------------------


def _norm_label(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def _sipri_numeric(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = re.sub(r"\[[^\]]*\]", "", str(value).strip()).replace("%", "").replace(" ", "").replace(",", "")
    if text.lower() in {".", "..", "...", "—", "-", "na", "n/a", "n a"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _sipri_year_columns(rows: list[list[Any]], max_header_rows: int = 12) -> dict[int, str]:
    years = {}
    for row in rows[:max_header_rows]:
        for idx, value in enumerate(row):
            m = re.fullmatch(r"(?:FY[- ]?)?(19\d{2}|20\d{2})", str(value or "").strip(), re.I)
            if m:
                years[idx] = m.group(1)
    return years


def _sipri_country_columns(rows: list[list[Any]], header_end: int) -> list[int]:
    exact = {"country", "countries", "country name", "country or area", "country area", "country/area", "economy", "economies", "name", "state", "area", "reference area"}
    scores = {}
    width = max((len(r) for r in rows), default=0)
    for row in rows[max(0, header_end - 5):header_end + 1]:
        for idx, value in enumerate(row):
            n = _norm_label(value)
            if n in exact:
                scores[idx] = scores.get(idx, 0) + 20
            elif "country" in n or "econom" in n or n == "area":
                scores[idx] = scores.get(idx, 0) + 7
    for idx in range(min(width, 10)):
        sample = [r[idx] for r in rows[header_end:header_end + 40] if idx < len(r)]
        textish = sum(1 for v in sample if v not in (None, "") and _sipri_numeric(v) is None)
        if textish >= 8:
            scores[idx] = scores.get(idx, 0) + 1
    return [idx for idx, _ in sorted(scores.items(), key=lambda x: (-x[1], x[0]))]


def _sipri_sheet_score(title: str, field_name: str) -> int:
    n = _norm_label(title)
    if field_name == "military_share_government" and ("government" in n or "govt" in n): return 50
    if field_name == "military_burden_gdp" and ("gdp" in n or "gross domestic" in n): return 50
    if field_name == "military_expenditure_real" and ("constant" in n or "us dollar" in n or "usd" in n): return 50
    return 0


def _sipri_context_score(rows: list[list[Any]], header_end: int, field_name: str) -> int:
    context = " ".join(_norm_label(v) for r in rows[max(0, header_end - 6):header_end + 1] for v in r)
    terms = {"military_expenditure_real": ("constant", "2024", "us"), "military_burden_gdp": ("share", "gross domestic product"), "military_share_government": ("share", "government expenditure")}[field_name]
    score = 0
    for term in terms:
        if _norm_label(term) in context or term in context:
            score += 10
    return score


def _sipri_parse_strategy(rows: list[list[Any]], title: str, field_name: str, strategy: int) -> dict[str, Any]:
    if not 1 <= strategy <= 20:
        raise ValueError(strategy)
    max_header = 18 if strategy % 2 == 0 else 12
    best = None
    for header_end in range(1, min(len(rows), max_header) + 1):
        years = _sipri_year_columns(rows[max(0, header_end - 6):header_end + 1], 7)
        if len(years) < 5:
            continue
        context_score = _sipri_context_score(rows, header_end, field_name)
        sheet_score = _sipri_sheet_score(title, field_name)
        if strategy <= 10 and context_score < 6 and sheet_score < 20:
            continue
        if 11 <= strategy <= 15 and context_score < 10 and sheet_score < 20:
            continue
        country_cols = _sipri_country_columns(rows, header_end)
        if not country_cols:
            continue
        if strategy in (3, 8, 13, 18): country_cols = [c for c in country_cols if c <= 1]
        elif strategy in (4, 9, 14, 19): country_cols = [c for c in country_cols if c <= 3]
        elif strategy in (5, 10, 15, 20): country_cols = country_cols[:1]
        for country_idx in country_cols:
            observations = {}
            for row in rows[header_end:]:
                if country_idx >= len(row): continue
                country = str(row[country_idx] or "").strip()
                if not country or _sipri_numeric(country) is not None: continue
                obs = []
                for col_idx, year in years.items():
                    if col_idx < len(row):
                        value = _sipri_numeric(row[col_idx])
                        if value is not None: obs.append({"period": str(year), "value": value})
                if obs: observations[country] = obs
            count = sum(len(v) for v in observations.values())
            if count < 20: continue
            score = context_score + sheet_score + min(count, 500) / 100.0 + min(header_end, 10) * 0.1
            candidate = {"strategy": strategy, "sheet": title, "header_end": header_end, "country_column": country_idx, "data": observations, "score": score}
            if best is None or score > best["score"]: best = candidate
    if best is None: raise RuntimeError(f"strategy {strategy}: no valid table")
    return best


def _find_sipri_tables(path: str) -> dict[str, dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    results = {}
    fields = ("military_expenditure_real", "military_burden_gdp", "military_share_government")
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows: continue
        for field_name in fields:
            for strategy in range(1, 21):
                try:
                    candidate = _sipri_parse_strategy(rows, ws.title, field_name, strategy)
                    current = results.get(field_name)
                    if current is None or candidate["score"] > current["score"]:
                        results[field_name] = candidate
                except Exception:
                    continue
    wb.close()
    return results


def update_sipri(cache: dict[str, Any]) -> None:
    path = None
    try:
        path = _download_temp(SIPRI_XLSX_URL, ".xlsx")
        tables = _find_sipri_tables(path)
        ts = now_iso()
        for field_name in ("military_expenditure_real", "military_burden_gdp", "military_share_government"):
            table = tables.get(field_name)
            if not table: raise RuntimeError(f"SIPRI: keine der 20 Parserstrategien fand {field_name}")
            grouped = table["data"]
            field = cache["SIPRI_DEFENCE"][field_name]
            field.update({"version": "2025-revised-2026-04-27", "status": "REAL", "retrieved_at": ts, "last_successful_update": ts, "data_period": max((obs["period"] for rows in grouped.values() for obs in rows), default=None), "data": {country: {"reference_area": country, "observations": obs} for country, obs in grouped.items()}, "series_count": len(grouped), "observation_count": sum(len(v) for v in grouped.values()), "source_notes": ["Offizielle SIPRI Military Expenditure Database, revidierte Fassung 27.04.2026.", f"Beste von 20 Parserstrategien: {table['strategy']} auf Blatt {table['sheet']} (Headerzeile {table['header_end']}, Länder-Spalte {table['country_column']}).", "Calendar-year basis for constant USD/GDP; government share follows financial year.", "Quelleneigene Schätzungen/Flags werden nicht als Python-Schätzungen erzeugt."]})
            LOG.info("SIPRI %s: beste Parserstrategie %d auf Blatt %s (Score %.1f)", field_name, table["strategy"], table["sheet"], table["score"])
    except Exception as exc:
        LOG.warning("SIPRI: %s", exc)
        for field_name in ("military_expenditure_real", "military_burden_gdp", "military_share_government"):
            field = cache["SIPRI_DEFENCE"][field_name]
            cache["SIPRI_DEFENCE"][field_name] = set_real_cached_if_valid(field, "SIPRI_DEFENCE")
    finally:
        if path:
            try: Path(path).unlink(missing_ok=True)
            except OSError: pass


# ---------------------------------------------------------------------------
# Prüfung / Status
# ---------------------------------------------------------------------------

def check_cache(cache: dict[str, Any]) -> tuple[bool, list[str]]:
    errors: list[str] = []

    expected = {
        "OECD_STAN": {
            "value_added", "investment", "labour_input"
        },
        "OECD_PRODUCTIVITY": {
            "labour_productivity_level", "labour_productivity_growth"
        },
        "AI_INDEX": {
            "ai_investment", "ai_adoption", "ai_compute"
        },
        "IEA_ELECTRICITY": {
            "electricity_generation", "electricity_balance"
        },
        "SIPRI_DEFENCE": {
            "military_expenditure_real",
            "military_burden_gdp",
            "military_share_government",
        },
    }

    for section, fields in expected.items():
        if section not in cache:
            errors.append(f"Fehlender Abschnitt: {section}")
            continue
        for field in fields:
            if field not in cache[section]:
                errors.append(f"Fehlendes Feld: {section}.{field}")

    if cache.get("cache_version") != CACHE_VERSION:
        errors.append(
            f"Unerwartete Cache-Version: {cache.get('cache_version')}"
        )

    return not errors, errors


def print_summary(cache: dict[str, Any]) -> None:
    print("NEUBER MACRO & MARKETS – STRUKTUR-TRENDS CACHE")
    print("=" * 56)
    print(f"Cache-Version : {cache.get('cache_version')}")
    print(f"Aktualisiert  : {cache.get('cache_updated_at')}")
    print()

    for section, fields in cache.items():
        if not isinstance(fields, dict) or section in (
            "cache_version", "cache_created_at", "cache_updated_at"
        ):
            continue

        print(f"[{section}]")
        for name, field in fields.items():
            summary = cache_field_summary(field)
            print(
                f"  {name:32s} "
                f"{summary['status']:14s} "
                f"period={str(summary['data_period']):12s} "
                f"obs={summary['observations']}"
            )
        print()


# ---------------------------------------------------------------------------
# Orchestrierung
# ---------------------------------------------------------------------------

def run(update: bool = True, start_period: int = 2000) -> int:
    cache = read_json(CACHE_FILE)

    # Struktur ergänzen, falls ein älterer/teilweiser Cache vorliegt.
    template = new_cache()
    for section, fields in template.items():
        if section in ("cache_version", "cache_created_at", "cache_updated_at"):
            continue
        cache.setdefault(section, {})
        for field_name, field in fields.items():
            cache[section].setdefault(field_name, field)

    if update:
        update_oecd_stan(cache, start_period=start_period)
        update_oecd_productivity(cache, start_period=start_period)
        update_ai_index(cache)
        update_iea(cache)
        update_sipri(cache)

    # Cache-Metadaten beim Übergang von einer älteren Cache-Version
    # kontrolliert auf die aktuelle Strukturversion migrieren.
    # Die vorhandenen Beobachtungen werden dabei nicht verworfen oder verändert.
    cache["cache_version"] = CACHE_VERSION
    cache["cache_updated_at"] = now_iso()
    atomic_write_json(CACHE_FILE, cache)

    ok, errors = check_cache(cache)
    print_summary(cache)

    if not ok:
        print("CHECK: FEHLER")
        for error in errors:
            print(f"  - {error}")
        return 2

    print("CHECK: OK")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Zentraler Cache für strukturelle Markttrends."
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="Nur Cache-Struktur prüfen, keine Quellen abrufen.",
    )
    parser.add_argument(
        "--show",
        action="store_true",
        help="Cache-Zusammenfassung anzeigen.",
    )
    parser.add_argument(
        "--start-period",
        type=int,
        default=2000,
        help="Ältestes abzurufendes Jahr (Standard: 2000).",
    )
    args = parser.parse_args()

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )

    if args.check or args.show:
        cache = read_json(CACHE_FILE)
        if args.show:
            print_summary(cache)
        ok, errors = check_cache(cache)
        if not ok:
            for error in errors:
                print(f"ERROR: {error}")
            return 2
        return 0

    return run(update=True, start_period=args.start_period)


if __name__ == "__main__":
    sys.exit(main())

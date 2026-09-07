#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Deterministischer Offline-Selbsttest für optionsschein_kurse.py.

Der Test führt KEINEN Internetabruf durch. Die Kursquelle wird kontrolliert
simuliert. Damit lassen sich XLSX-Schema, Zahlentypen, Formate und die
Trennung zwischen OS- und Nicht-OS-Zeilen reproduzierbar prüfen.
"""

from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import optionsschein_kurse as osk


BASE = Path(__file__).resolve().parent
FIXTURE = BASE / "Offene_Positionen_TEST.xlsx"

QUOTES = {
    "UN68EW": 0.40,
    "PM0216": 1.27,
    "PM026C": 0.80,
    "PK9UHB": 1.60,
    "GW5GT1": 1.90,
    "UN4N9T": 0.50,
    "HM0ZHD": 4.80,
}


class OptionsscheinXlsxTest(unittest.TestCase):
    def test_parse_fallback_last_price(self):
        html = (
            "WKN PM0216 Geld - Brief - "
            "Daten & Zahlen Kursdaten Letzter Preis 1,27 G"
        )
        quote = osk.parse_stuttgart_html(html, "PM0216")
        self.assertEqual(quote.aktueller_kurs, 1.27)

    def test_parse_prefers_geld(self):
        html = "WKN UG9RWH Geld 1,46 Brief 1,47 Letzter Preis 1,45 G"
        quote = osk.parse_stuttgart_html(html, "UG9RWH")
        self.assertEqual(quote.aktueller_kurs, 1.46)

    def test_parse_rejects_wrong_instrument(self):
        html = "WKN ABCDEF Geld 1,00 Brief 1,01"
        with self.assertRaises(osk.OptionsscheinAbrufFehler):
            osk.parse_stuttgart_html(html, "UN68EW")

    def test_xlsx_update(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "Offene_Positionen_TEST.xlsx"
            shutil.copy2(FIXTURE, target)

            def fake_quote(wkn: str):
                return osk.OptionsscheinKurs(wkn=wkn, aktueller_kurs=QUOTES[wkn])

            with patch.object(osk, "hole_optionsschein_kurs", side_effect=fake_quote):
                count = osk.verarbeite_datei(target)

            self.assertEqual(count, 7)

            wb = load_workbook(target, data_only=False)
            ws = wb["Offene_Positionen"]
            headers = {c.value: c.column for c in ws[1]}

            self.assertEqual(
                [ws.cell(1, headers[x]).value for x in osk.OS_SPALTEN],
                osk.OS_SPALTEN,
            )

            for row in range(2, ws.max_row + 1):
                if str(ws.cell(row, headers["Produkt_Typ"]).value or "").strip().lower() != "optionsschein":
                    continue
                wkn = ws.cell(row, headers["OS_WKN"]).value
                current = ws.cell(row, headers["OS_Aktueller_Kurs"]).value
                entry = ws.cell(row, headers["OS_Einstiegskurs"]).value
                performance = ws.cell(row, headers["OS_Performance%"]).value

                self.assertIsInstance(current, (int, float))
                self.assertAlmostEqual(current, QUOTES[wkn], places=10)
                self.assertAlmostEqual(
                    performance,
                    (QUOTES[wkn] - entry) / entry * 100.0,
                    places=10,
                )
                self.assertEqual(ws.cell(row, headers["OS_Quelle"]).value, "Börse Stuttgart")
                self.assertTrue(ws.cell(row, headers["OS_Zeitstempel"]).value)
                self.assertEqual(ws.cell(row, headers["OS_Aktueller_Kurs"]).number_format, "0.00")
                self.assertEqual(ws.cell(row, headers["OS_Performance%"]).number_format, '0.00" %"')
                self.assertEqual(ws.cell(row, headers["OS_WKN"]).number_format, "@")

            wb.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)

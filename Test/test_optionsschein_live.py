#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Live-Integrationstest für die Optionsschein-Kursaktualisierung.

Dieser Test ruft die sieben echten WKNs aus der Test-XLSX bei der Börse Stuttgart
ab und schreibt ausschließlich in eine temporäre Kopie der Test-XLSX.
Die Produktivdatei wird weder gelesen noch verändert.
"""

from __future__ import annotations

import hashlib
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import optionsschein_kurse as osk


BASE = Path(__file__).resolve().parent
FIXTURE = BASE / "Offene_Positionen_TEST.xlsx"
EXPECTED_WKNS = {
    "UN68EW",
    "PM0216",
    "PM026C",
    "PK9UHB",
    "GW5GT1",
    "UN4N9T",
    "HM0ZHD",
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


class OptionsscheinLiveIntegrationTest(unittest.TestCase):
    def test_live_stuttgart_to_xlsx(self):
        if not FIXTURE.exists():
            self.fail(f"Testdatei fehlt: {FIXTURE}")

        fixture_before = sha256(FIXTURE)

        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "Offene_Positionen_TEST_LIVE.xlsx"
            shutil.copy2(FIXTURE, target)

            count = osk.verarbeite_datei(target)
            self.assertEqual(count, len(EXPECTED_WKNS))

            wb = load_workbook(target, data_only=False)
            self.assertIn(osk.SHEET_NAME, wb.sheetnames)
            ws = wb[osk.SHEET_NAME]
            headers = {c.value: c.column for c in ws[1]}

            seen = set()
            for row in range(2, ws.max_row + 1):
                produkt_typ = str(
                    ws.cell(row, headers["Produkt_Typ"]).value or ""
                ).strip().lower()
                if produkt_typ != "optionsschein":
                    continue

                wkn = str(ws.cell(row, headers["OS_WKN"]).value or "").strip().upper()
                seen.add(wkn)
                current = ws.cell(row, headers["OS_Aktueller_Kurs"]).value
                performance = ws.cell(row, headers["OS_Performance%"]).value
                source = ws.cell(row, headers["OS_Quelle"]).value
                timestamp = ws.cell(row, headers["OS_Zeitstempel"]).value

                self.assertIn(wkn, EXPECTED_WKNS)
                self.assertIsInstance(current, (int, float), f"{wkn}: Kurs fehlt")
                self.assertGreaterEqual(current, 0, f"{wkn}: Kurs negativ")
                self.assertEqual(source, osk.QUELLE, f"{wkn}: falsche Quelle")
                self.assertTrue(timestamp, f"{wkn}: Zeitstempel fehlt")
                self.assertIsInstance(performance, (int, float), f"{wkn}: Performance fehlt")
                self.assertEqual(
                    ws.cell(row, headers["OS_Aktueller_Kurs"]).number_format,
                    osk.OS_NUMMER_FORMAT,
                )
                self.assertEqual(
                    ws.cell(row, headers["OS_Performance%"]).number_format,
                    osk.OS_PERFORMANCE_FORMAT,
                )
                self.assertEqual(
                    ws.cell(row, headers["OS_WKN"]).number_format,
                    "@",
                )

                print(
                    f"[LIVE OK] {wkn}: Kurs={current}, "
                    f"Performance={performance:.4f}%, Quelle={source}, Zeit={timestamp}"
                )

            wb.close()

            self.assertEqual(seen, EXPECTED_WKNS)

        # Die Test-Fixture selbst darf durch den Live-Test nicht verändert werden.
        self.assertEqual(
            sha256(FIXTURE),
            fixture_before,
            "Die Test-Fixture wurde unerwartet verändert.",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)
